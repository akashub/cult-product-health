"""Builds a small synthetic workbook matching config.example.yaml, with the
same kinds of mess the real sheet has (float IDs, backtick codes, header
spaces, spelling variants, blank rows). No real data."""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

import openpyxl

MONTHS = {1: "1~January", 2: "2~February", 3: "3~March"}


def build(path: Path) -> dict:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    t = wb.create_sheet("Tickets")
    t.append(["Ticket Id", "Created", "Month", "Platform", "Order ID", "Model", "L1", "L2", "L3", "Subject", "Email"])
    t.append([1001.0, datetime(2026, 1, 3, 10), MONTHS[1], "Amazon", "403-1234567-1234567", "Gun A", "Product", "Product", "Dead", "x", "a@b.c"])
    t.append([1002.0, datetime(2026, 1, 9, 11), MONTHS[1], "No details", "OD100000009", "Not Mentioned", "Product", "Refund", "Late", "y", "a@b.c"])
    t.append([1003.0, datetime(2026, 2, 1, 9), MONTHS[2], "1P", 100001.0, "foot b", "Product", "Product", "Noise", "z", "a@b.c"])
    t.append([None] * 11)  # blank row must be ignored, not counted

    a = wb.create_sheet("Approved")
    a.append(["Date", "Month", "Order ID", "Batch", "Model", "SKU ", "Amount ", "Issue", "Mode", "Channel", "AWB", "Warehouse"])
    rows = [
        (datetime(2026, 1, 5), 1, "405-1111111-2222222", "B1", "Gun A", "SKU-A1", 1999.0, "Dead", "Exchange", "3P", 9000000000004.0, "W1"),
        (datetime(2026, 1, 6), 1, 4000000001.0, "B1", "Sample Gun A Blue", "SKU-A1", 1799.0, "Noise", "Return", "3P", "111  222", "W1"),
        (datetime(2026, 2, 7), 2, "OD100000000000000002", "NA", "Foot B", "SKU-B2", 2199.0, "Dead", "exchange", "1P", None, "W2"),
        (datetime(2026, 3, 8), 3, 100002.0, "B2", "Gun A", "SKU-B1", 999.0, "Heat", "Return", "1P", None, "W2"),  # SKU/name conflict
    ]
    for d, m, *rest in rows:
        a.append([d, MONTHS[m], *rest])

    p = wb.create_sheet("Pending")
    p.append(["Order ID", "SKU", "Model", "Created", "Month", "Reason", "Created By", "Comment"])
    p.append([200003.0, "SKU-A1", "Gun A", datetime(2026, 1, 16), MONTHS[1], "Quality is poor", "CUSTOMER", None])
    p.append([200004.0, "SKU-B1", "Foot B", datetime(2026, 3, 16), MONTHS[3], "Damaged", "CUSTOMER", "not charging"])
    p.append([None] * 8)

    w = wb.create_sheet("Warehouse")
    w.append(["Created", "Month", "Channel", "Channel Type", "Return Type", "Return Sub Type", "Reason", "Order ID",
              "Item Code", "SKU", "Model", "Facility", "Status"])
    w.append([datetime(2026, 2, 23), MONTHS[2], "STORE-SHOP", "1P", "RVP", "Customer Return", "Quality is poor",
              200001.0, "`30000000000001", "SKU-A1", "Gun A", "F1", "COMPLETE"])
    w.append([datetime(2026, 3, 27), MONTHS[3], "MARKET-X", "3P", "RTO", "Courier Return", None,
              200002.0, "`30000000000002", "SKU-B2", "Foot B", "F2", "RETURNED"])

    s = wb.create_sheet("Sales")
    s.append(["Month", "SKU", "Model", "Platform", "Units", "Revenue"])
    for m, sku, units in [(1, "SKU-A1", 100), (1, "SKU-B1", 50), (2, "SKU-B2", 40), (3, "SKU-A1", 20), (3, "SKU-B1", 10)]:
        s.append([datetime(2026, m, 1), sku, None, "Amazon", units, units * 1999.0])

    # Hand-made pivot the judge must reproduce: approved per month, total = approved + pending
    d = wb.create_sheet("Dashboard")
    d["B1"], d["C1"] = "Approved", "Total"
    expected = {1: (2, 3), 2: (1, 1), 3: (1, 2)}
    for i, m in enumerate((1, 2, 3), start=2):
        d[f"B{i}"], d[f"C{i}"] = expected[m]
    d["B5"], d["C5"] = sum(v[0] for v in expected.values()), sum(v[1] for v in expected.values())

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return expected


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "data" / "sample.xlsx"
    build(out)
    print(f"wrote {out}")
