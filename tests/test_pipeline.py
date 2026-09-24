"""End-to-end on the synthetic workbook: parse -> judge -> reconcile."""

from pathlib import Path

import pytest

from cultph.config import ROOT, load_config
from cultph.ingest import ingest
from cultph.judge import run_checks, verdict
from cultph.metrics import breakdown, dashboard_recompute, rows_for
from cultph.sources import XlsxSource

from make_fixture import build


@pytest.fixture()
def run(tmp_path: Path):
    wb = tmp_path / "sample.xlsx"
    build(wb)
    cfg = load_config(ROOT / "config.example.yaml")
    src = XlsxSource(wb)
    res = ingest(src, cfg)
    return cfg, src, res


def test_rows_conserved_and_blank_rows_ignored(run):
    _, _, res = run
    assert res.source_rows == {"tickets": 3, "approved": 4, "pending": 2, "wms": 2}
    assert not res.rejects and not res.header_errors


def test_ids_are_clean_strings(run):
    _, _, res = run
    appr = res.tables["approved"]
    assert list(appr["order_id"]) == ["405-1111111-2222222", "4000000001", "OD100000000000000002", "100002"]
    assert appr.loc[1, "awbs"] == "111,222"
    assert res.tables["wms"].loc[0, "item_code"] == "30000000000001"


def test_product_resolution(run):
    _, _, res = run
    appr = res.tables["approved"]
    assert list(appr["product"]) == ["Sample Gun A", "Sample Gun A", "Sample Foot B", "Sample Foot B"]
    assert list(appr["name_conflict"]) == [False, False, False, True]  # last row: SKU says B, name says A
    tix = res.tables["tickets"]
    assert list(tix["map_method"]) == ["alias", "not_mentioned", "alias"]
    assert list(tix["platform"]) == ["Amazon", "Flipkart", "Cult Store"]
    assert list(tix["platform_inferred"]) == [False, True, False]
    assert list(tix["is_product_issue"]) == [True, False, True]


def test_mode_normalised(run):
    _, _, res = run
    assert set(res.tables["approved"]["mode"]) == {"Exchange", "Return"}


def test_dashboard_reconciles(run):
    cfg, src, res = run
    got = dashboard_recompute(res.tables)
    assert got.to_dict("records") == [{"month": "2026-01", "approved": 2, "total": 3},
                                      {"month": "2026-02", "approved": 1, "total": 1},
                                      {"month": "2026-03", "approved": 1, "total": 2}]
    checks = {c["check"]: c for c in run_checks(res, cfg, src)}
    assert checks["dashboard_reconcile"]["status"] == "pass"
    assert checks["sku_name_conflict:approved"]["status"] == "warn"
    assert verdict(list(checks.values())) == "warn"


def test_judge_blocks_on_dashboard_mismatch(run, tmp_path):
    cfg, _, res = run
    import openpyxl
    wb_path = tmp_path / "bad.xlsx"
    build(wb_path)
    wb = openpyxl.load_workbook(wb_path)
    wb["Dashboard"]["B2"] = 99
    wb.save(wb_path)
    checks = run_checks(res, cfg, XlsxSource(wb_path))
    assert verdict(checks) == "fail"


def test_judge_blocks_when_sheet_has_extra_month(run, tmp_path):
    """A new month in the sheet's pivot that the raw data lacks must not pass silently."""
    cfg, _, res = run
    import openpyxl
    wb_path = tmp_path / "extra.xlsx"
    build(wb_path)
    wb = openpyxl.load_workbook(wb_path)
    d = wb["Dashboard"]
    d["B5"], d["C5"], d["B6"], d["C6"] = 5, 5, 9, 11  # extra month row + new grand total
    wb.save(wb_path)
    checks = {c["check"]: c for c in run_checks(res, cfg, XlsxSource(wb_path))}
    assert checks["dashboard_reconcile"]["status"] == "fail"
    assert "month rows" in checks["dashboard_reconcile"]["detail"]


def test_judge_blocks_on_missing_header(tmp_path):
    wb = tmp_path / "s.xlsx"
    build(wb)
    cfg = load_config(ROOT / "config.example.yaml")
    cfg.raw["tabs"]["approved"]["columns"]["issue"] = "Issue Typo"
    res = ingest(XlsxSource(wb), cfg)
    assert verdict(run_checks(res, cfg)) == "fail"


def test_breakdown_and_drilldown_agree(run):
    _, _, res = run
    appr = res.tables["approved"]
    b = breakdown(appr, "issue", within="product")
    assert (b.groupby("product")["share"].sum() - 1).abs().max() < 1e-12
    for r in b.itertuples():
        assert len(rows_for(appr, {"product": r.product, "issue": r.issue})) == r.count


def test_alias_collision_rejected(tmp_path):
    cfg = load_config(ROOT / "config.example.yaml")
    cfg.products[1].aliases.append("gun-a")  # same key as product A's alias "Gun A"
    with pytest.raises(ValueError):
        cfg.alias_index()
