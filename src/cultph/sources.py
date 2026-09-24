"""Readers that turn a workbook (local .xlsx or a Google Sheet) into the same
shape: {tab_name: (headers, rows)} with raw cell values. Everything downstream
is source-agnostic."""

from __future__ import annotations

import warnings
from pathlib import Path

from .config import DATA_DIR

Table = tuple[list[str], list[list]]

SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets.readonly",
    "https://www.googleapis.com/auth/drive.metadata.readonly",
]


class XlsxSource:
    def __init__(self, path: str | Path):
        import openpyxl

        self.path = Path(path)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # pivot-cache warnings are noise
            self._wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)

    def table(self, tab: str) -> Table:
        rows = [list(r) for r in self._wb[tab].iter_rows(values_only=True)]
        if not rows:
            return [], []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        return headers, rows[1:]

    def cells(self, tab: str, cell_range: str) -> list:
        ws = self._wb[tab]
        out = []
        for row in ws[cell_range]:
            for c in row if isinstance(row, tuple) else (row,):
                out.append(c.value)
        return out


class GSheetSource:
    """Reads a Google Sheet the user can access, via desktop OAuth.

    Put the OAuth client JSON (type "Desktop app") at data/google_client.json;
    the first run opens a browser to sign in with Gmail and caches the token in
    data/google_token.json."""

    def __init__(self, spreadsheet_id: str):
        import gspread

        self._gc = gspread.oauth(
            scopes=SCOPES,
            credentials_filename=str(DATA_DIR / "google_client.json"),
            authorized_user_filename=str(DATA_DIR / "google_token.json"),
        )
        self._sh = self._gc.open_by_key(spreadsheet_id)

    def table(self, tab: str) -> Table:
        values = self._sh.worksheet(tab).get_all_values(value_render_option="UNFORMATTED_VALUE",
                                                        date_time_render_option="SERIAL_NUMBER")
        if not values:
            return [], []
        headers = [str(h).strip() for h in values[0]]
        return headers, [list(r) for r in values[1:]]

    def cells(self, tab: str, cell_range: str) -> list:
        vals = self._sh.worksheet(tab).get(cell_range, value_render_option="UNFORMATTED_VALUE")
        return [v for row in vals for v in (row or [None])]


def list_shared_sheets() -> list[dict]:
    """Spreadsheets visible to the signed-in Gmail account (owned or shared)."""
    import gspread

    gc = gspread.oauth(
        scopes=SCOPES,
        credentials_filename=str(DATA_DIR / "google_client.json"),
        authorized_user_filename=str(DATA_DIR / "google_token.json"),
    )
    return gc.list_spreadsheet_files()


def open_source(source_cfg: dict):
    kind = source_cfg.get("type", "xlsx")
    if kind == "xlsx":
        from .config import ROOT

        p = Path(source_cfg["path"])
        return XlsxSource(p if p.is_absolute() else ROOT / p)
    if kind == "gsheet":
        return GSheetSource(source_cfg["spreadsheet_id"])
    raise ValueError(f"unknown source type {kind!r}")
